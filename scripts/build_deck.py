"""Generate docs/walkthrough-deck.pptx — assignment deliverable §5.6.

Run: ./.venv/bin/python scripts/build_deck.py
Requires python-pptx (a build-time tool only; not a runtime dependency of the
analyzer, so it is deliberately absent from pyproject.toml).

Content is grounded in README.md and docs/design.md. Keep them in sync: every
claim on a slide must be one the repository can actually support.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x16, 0x24, 0x3F)
NAVY_SOFT = RGBColor(0x24, 0x36, 0x5C)
STEEL = RGBColor(0x4A, 0x7F, 0xB5)
ICE = RGBColor(0xE8, 0xEE, 0xF7)
ICE_DEEP = RGBColor(0xD2, 0xDE, 0xEF)
AMBER = RGBColor(0xD9, 0x8A, 0x21)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5B, 0x6B, 0x85)
MUTED_LIGHT = RGBColor(0xA9, 0xB6, 0xCC)
GREEN = RGBColor(0x2C, 0x6E, 0x49)
RED = RGBColor(0x9B, 0x2C, 0x2C)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

W, H = 13.333, 7.5
M = 0.62
CW = W - M * 2

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return s


def shape(s, kind, x, y, w, h, fill=None, adj=None):
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    if adj is not None:
        try:
            sp.adjustments[0] = adj
        except (IndexError, ValueError):
            pass
    sp.text_frame.word_wrap = True
    return sp


def card(s, x, y, w, h, fill=ICE):
    return shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, adj=0.08)


def _spacing(run, hundredths):
    run.font._rPr.set("spc", str(int(hundredths)))


def _bullet(p, indent=0.17):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(indent)).emu))
    pPr.set("indent", str(-Emu(Inches(indent)).emu))
    for tag, attr, val in (("a:buFont", "typeface", "Arial"), ("a:buChar", "char", "•")):
        el = pPr.makeelement(qn(tag), {attr: val})
        pPr.append(el)


def text(s, body, x, y, w, h, *, size=13, bold=False, italic=False, color=NAVY,
         font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spc=None,
         bullets=False, space_after=0, line_spacing=None):
    """body: a string, or a list of strings rendered as separate paragraphs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    lines = body if isinstance(body, list) else [body]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        f = run.font
        f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
        f.color.rgb = color
        if spc:
            _spacing(run, spc)
        if bullets:
            _bullet(p)
    return tb


def numdot(s, x, y, label, d=0.42):
    shape(s, MSO_SHAPE.OVAL, x, y, d, d, AMBER)
    text(s, label, x, y, d, d, size=13, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def heading(s, title, kicker=None):
    if kicker:
        text(s, kicker.upper(), M, 0.42, CW, 0.26, size=11.5, bold=True,
             color=AMBER, spc=160)
    text(s, title, M, 0.7 if kicker else 0.5, CW, 0.66, size=31, bold=True,
         color=NAVY, font=HEAD_FONT)


def notes(s, body):
    s.notes_slide.notes_text_frame.text = body


# ------------------------------------------------------------- 1. title
s = slide(NAVY)
shape(s, MSO_SHAPE.OVAL, 10.6, -1.5, 4.6, 4.6, NAVY_SOFT)
shape(s, MSO_SHAPE.OVAL, 12.1, 5.3, 2.4, 2.4, NAVY_SOFT)
text(s, "CONTRACT ANALYZER", M, 1.95, 10.6, 0.32, size=12, bold=True,
     color=AMBER, spc=240)
text(s, ["Agentic RAG for", "contract compliance analysis"], M, 2.4, 10.4, 1.8,
     size=38, bold=True, color=WHITE, font=HEAD_FONT, line_spacing=1.08)
text(s, "Five fixed compliance questions · verbatim evidence · validated JSON · KPI dashboard",
     M, 4.36, 10.4, 0.36, size=15, italic=True, color=MUTED_LIGHT)
shape(s, MSO_SHAPE.RECTANGLE, M, 5.28, 1.1, 0.035, AMBER)
text(s, "Walkthrough deck — 5–10 minutes", M, 5.5, 6.0, 0.3, size=13, color=WHITE)
text(s, "Confidential — assignment material. Not for redistribution.",
     M, 5.86, 8.0, 0.28, size=11, color=MUTED_LIGHT)
notes(s, "Thirty-second frame: the deliverable is an end-to-end analyzer, but the "
         "interesting part is that every number and every quote it shows can be traced "
         "back to a source span, and anything it cannot measure it reports as "
         "unavailable rather than zero.")

# --------------------------------------------------------- 2. what it does
s = slide()
heading(s, "What the system does", "Overview")
y = 1.78
for n, head, body in [
    ("1", "Upload", "A PDF goes in through the Streamlit UI. Parsing, chunking and indexing are local."),
    ("2", "Analyse", "An agentic RAG loop answers five fixed compliance questions, question by question."),
    ("3", "Return", "Schema-validated JSON renders in the UI, alongside a KPI dashboard and audit trail."),
]:
    numdot(s, M, y + 0.04, n)
    text(s, head, M + 0.62, y, 6.0, 0.3, size=17, bold=True)
    text(s, body, M + 0.62, y + 0.34, 6.1, 0.7, size=13.5, color=MUTED)
    y += 1.34

card(s, 7.55, 1.72, 5.16, 4.0)
text(s, "Per-question output schema", 7.87, 1.96, 4.6, 0.3, size=14, bold=True)
text(s, [
    "Compliance Question",
    "Compliance State — Fully / Partially / Non-Compliant",
    "Confidence",
    "Relevant Quotes — verbatim, with section or exhibit reference",
    "Rationale",
], 7.9, 2.42, 4.5, 2.9, size=12.5, bullets=True, space_after=9)
text(s, "An unresolved question surfaces as a review item — never as a fourth state.",
     7.87, 5.84, 4.7, 0.5, size=11, italic=True, color=MUTED)
notes(s, "The three states are the only states. A question the pipeline could not resolve "
         "is escalated for review rather than quietly classified, because a wrong "
         "'Non-Compliant' and an honest 'we could not tell' have very different costs "
         "to a compliance team.")

# ------------------------------------------------------ 3. why this is hard
s = slide()
heading(s, "Why this document is hard", "Problem shape")
cw = (CW - 0.52) / 3
for i, (head, body) in enumerate([
    ("Answers live in tables",
     "Exhibits carry control rows the prose never states. One of the five questions is "
     "answerable only from a table row. Flattening tables to text destroys the row "
     "semantics that ground a citation."),
    ("Exceptions modify controls",
     "An open risk acceptance can qualify a control the main body grants. A shared "
     "identifier is not enough to downgrade a check — scope, dates and normative "
     "authority all have to hold."),
    ("There is no ground truth",
     "One sample contract ships, with zero labels. Any accuracy figure needs a hand-built "
     "gold set, so the system says 'unavailable' rather than imply a measurement it "
     "never made."),
]):
    x = M + i * (cw + 0.26)
    card(s, x, 1.78, cw, 3.55)
    numdot(s, x + 0.34, 2.1, str(i + 1), 0.4)
    text(s, head, x + 0.34, 2.7, cw - 0.68, 0.64, size=16, bold=True)
    text(s, body, x + 0.34, 3.34, cw - 0.68, 1.9, size=12.5, color=MUTED)

card(s, M, 5.62, CW, 0.84, NAVY)
text(s, "Consequence for the design: structure is discovered from the document at runtime. "
        "No section numbers, exhibit letters or ID prefixes are hand-transcribed into configuration.",
     M + 0.34, 5.62, CW - 0.68, 0.84, size=13.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "This slide explains every later decision. A mapping table transcribed from one PDF "
         "is an annotation of that PDF, not a system that analyses contracts — it fails "
         "silently the moment a contract numbers its clauses differently.")

# ---------------------------------------------------------- 4. architecture
s = slide()
heading(s, "Architecture", "Layers, and why each boundary exists")
y = 1.7
for i, (name, parts, why) in enumerate([
    ("Clients", "Streamlit UI · chat · MCP stdio",
     "Swappable. Replacing Streamlit with React changes nothing below it."),
    ("Service", "core/pipeline.py · core/tools/",
     "One analyze(), one ToolRegistry — a tool cannot behave differently in chat than in MCP."),
    ("Engine", "ingest · retrieval · rubric · agents",
     "The rubric is data (questions.yaml), so checks change without touching orchestration."),
    ("Persistence", "core/store/ — SQLite, WAL",
     "History and evidence snapshots render a past run without the original PDF."),
    ("Observability", "core/obs/ — JSONL + OpenTelemetry",
     "One span tree feeds terminal, file and Langfuse. No duplicate instrumentation."),
]):
    card(s, M, y, CW, 0.84, ICE if i % 2 == 0 else ICE_DEEP)
    text(s, name, M + 0.34, y, 1.9, 0.84, size=15, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, parts, M + 2.24, y, 3.5, 0.84, size=12, color=STEEL, anchor=MSO_ANCHOR.MIDDLE)
    text(s, why, M + 5.85, y, CW - 6.2, 0.84, size=12.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.96
notes(s, "The load-bearing boundary is the service layer. mcp_server.py imports the same "
         "service layer as the UI, which keeps the MCP deliverable executable rather than "
         "diagrammatic. No business logic lives in a Streamlit callback.")

# -------------------------------------------------------- 5. the agent loop
s = slide()
heading(s, "Agentic design — the bounded loop", "§3.2 required")
nodes = ["Router", "Retrieve", "Judge", "Verify", "Critic"]
bw, bh, gap, x0 = 1.92, 0.86, 0.44, M + 0.1
for i, n in enumerate(nodes):
    x = x0 + i * (bw + gap)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.95, bw, bh,
          AMBER if i == 3 else NAVY, adj=0.12)
    text(s, n, x, 1.95, bw, bh, size=15, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(nodes) - 1:
        shape(s, MSO_SHAPE.RIGHT_ARROW, x + bw + 0.09, 2.27, 0.26, 0.22, STEEL)

vx = x0 + 3 * (bw + gap)
shape(s, MSO_SHAPE.DOWN_ARROW, vx + bw / 2 - 0.11, 2.86, 0.22, 0.26, AMBER)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, vx - 0.24, 3.18, bw + 0.48, 0.62, ICE_DEEP, adj=0.16)
text(s, "Quote repair ×1", vx - 0.24, 3.18, bw + 0.48, 0.62, size=12.5, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, "critic requests a specific follow-up → back to Router (bounded budget)",
     M, 4.1, CW, 0.3, size=12, italic=True, color=STEEL, align=PP_ALIGN.CENTER)

y = 4.64
for k, v in [
    ("Framework", "LangGraph. Explicit nodes and conditional edges make routing and stopping "
                  "inspectable state, not prompt behaviour."),
    ("Protocol", "Typed Python records and validated JSON between agents — never free-form "
                 "inter-agent chat."),
    ("Failure", "Per-question budgets on calls, tokens and wall time. Verified findings survive "
                "an early stop; the rest degrade to UNRESOLVED, never to a guess."),
]:
    text(s, k, M, y, 1.5, 0.5, size=13, bold=True, color=AMBER)
    text(s, v, M + 1.55, y, CW - 1.55, 0.54, size=12.5, color=MUTED)
    y += 0.62
notes(s, "Walk the loop once. Router plans, retrieve gathers a shared evidence pack, judge "
         "assesses each required facet, verify is mechanical, critic is semantic. The amber "
         "node is the only place a quotation error gets a second chance, capped at one "
         "repair per check.")

# -------------------------------------------------------------- 6. router
s = slide()
heading(s, "Router agent", "Decompose, then dispatch")
card(s, M, 1.72, 5.6, 4.05)
text(s, "Decomposition", M + 0.34, 1.98, 4.9, 0.3, size=15, bold=True)
text(s, [
    "One compliance question becomes several sub-requirement checks.",
    "Each check is planned and scored on its own, so a grouped question cannot pass on one covered facet.",
    "Sub-requirements come from questions.yaml — data, not code.",
], M + 0.38, 2.44, 4.85, 2.3, size=13, color=MUTED, bullets=True, space_after=10)

card(s, M + 5.92, 1.72, CW - 5.92, 4.05)
text(s, "Dispatch", M + 6.26, 1.98, 5.0, 0.3, size=15, bold=True)
ty = 2.48
for n, d in [
    ("Clause search", "lexical + dense retrieval over structural chunks"),
    ("Table fetch", "whole tables and rows, with headers preserved"),
    ("Cross-reference", "follow a reference to the clause it names"),
    ("Modifier scan", "document-wide sweep for exceptions and waivers"),
]:
    shape(s, MSO_SHAPE.OVAL, M + 6.26, ty + 0.07, 0.16, 0.16, AMBER)
    text(s, n, M + 6.56, ty, 2.4, 0.28, size=13, bold=True)
    text(s, d, M + 6.56, ty + 0.29, 4.6, 0.46, size=11.5, color=MUTED)
    ty += 0.82

card(s, M, 5.95, CW, 0.74, NAVY)
text(s, "No hand-written section map. Where evidence lives is discovered per document — "
        "that is what makes the router an agent rather than a lookup table.",
     M + 0.34, 5.95, CW - 0.68, 0.74, size=13, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "If asked why not one retrieval call: a flat query cannot satisfy a seven-part "
         "password requirement, and gives you nothing to attribute a partial failure to. "
         "Decomposition is also what makes the deterministic aggregation defensible.")

# ----------------------------------------------------------- 7. evaluator
s = slide()
heading(s, "Evaluator agent", "Mechanical verifier + semantic critic")
card(s, M, 1.72, 6.0, 3.9)
text(s, "Mechanical verifier — deterministic", M + 0.34, 1.98, 5.3, 0.3, size=15, bold=True)
text(s, [
    "schema — verdict, confidence and fields are well formed",
    "source_in_pack — every citation names evidence actually supplied",
    "quote_verbatim — the quote is a literal span of its chunk",
    "quote_in_document — and of the document as a whole",
    "not_found_consistency — NOT_FOUND carries no quote",
    "modifier_citations — an applied exception must cite its source",
], M + 0.38, 2.44, 5.25, 2.95, size=12.5, color=MUTED, bullets=True, space_after=7)

card(s, M + 6.32, 1.72, CW - 6.32, 3.9)
text(s, "Semantic critic — judgement", M + 6.66, 1.98, 5.0, 0.3, size=15, bold=True)
cy = 2.48
for k, v in [
    ("accept", "evidence entails the verdict; confidence is capped at the critic's own"),
    ("retrieve", "a specific, actionable retrieval objective — loops back to the router"),
    ("unresolved", "a named evidence gap, reported rather than guessed"),
]:
    text(s, k, M + 6.66, cy, 1.6, 0.28, size=13.5, bold=True, color=AMBER)
    text(s, v, M + 6.66, cy + 0.3, 5.1, 0.64, size=12.5, color=MUTED)
    cy += 1.05
text(s, "It also checks scope, dates, normative authority and contradictions.",
     M + 6.66, 5.08, 5.1, 0.4, size=11.5, italic=True, color=STEEL)

card(s, M, 5.82, CW, 0.84, NAVY)
text(s, "Hallucination detection is the verbatim source check, not a vibe. A semantic critic "
        "can never overrule mechanical verification — it can only ask for more evidence, or "
        "declare the check unresolved.",
     M + 0.34, 5.82, CW - 0.68, 0.84, size=13, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Confidence calibration: the critic's confidence caps the judge's, never raises it, "
         "and the dashboard bins predicted confidence against observed correctness wherever "
         "reviewed labels exist. Where they do not, the bin is absent rather than zero.")

# --------------------------------------------------------- 8. aggregation
s = slide()
heading(s, "Deterministic aggregation", "Sub-requirement verdicts → one compliance state")
y = 1.8
for state, when, colour in [
    ("Fully Compliant", "every required sub-requirement is MET on verified evidence", GREEN),
    ("Partially Compliant", "some MET, some missing, contradicted or narrowed by an exception", AMBER),
    ("Non-Compliant", "the obligation is absent or contradicted across the document", RED),
]:
    card(s, M, y, 7.4, 0.94)
    shape(s, MSO_SHAPE.OVAL, M + 0.34, y + 0.37, 0.2, 0.2, colour)
    text(s, state, M + 0.68, y, 2.5, 0.94, size=14.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, when, M + 3.1, y, 4.1, 0.94, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    y += 1.06

card(s, M + 7.72, 1.8, CW - 7.72, 3.0, NAVY)
text(s, "Why a rule, not a second model call", M + 8.06, 2.04, 3.7, 0.32,
     size=14, bold=True, color=WHITE)
text(s, [
    "The same evidence always yields the same state.",
    "A disagreement is traceable to one named sub-requirement.",
    "The rule is readable by a compliance reviewer who does not read Python.",
], M + 8.1, 2.46, 3.63, 2.2, size=12, color=ICE, bullets=True, space_after=9)

card(s, M, 5.18, CW, 1.3, ICE_DEEP)
text(s, "Citations are rendered from the chunk the quote came from",
     M + 0.34, 5.36, 11.0, 0.3, size=13.5, bold=True)
text(s, "“Section 6.6” · “Exhibit G (GOV-04)” · “Page 21” — derived from the "
        "document's own structure, and stored per run so a historical result still cites "
        "correctly once the uploaded PDF is gone.",
     M + 0.34, 5.7, CW - 0.68, 0.62, size=12.5, color=MUTED)
notes(s, "Standing preference behind this slide: no single-prompt classification. The model "
         "judges facets against evidence; the state is arithmetic over those judgements.")

# ------------------------------------------------------ 9. connector / MCP
s = slide()
heading(s, "Connector / MCP design", "§3.3 required")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, M + 3.9, 1.78, 4.4, 0.82, NAVY, adj=0.12)
text(s, "One ToolRegistry  (core/tools/)", M + 3.9, 1.78, 4.4, 0.82, size=14, bold=True,
     color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
sw = 3.6
for i, (n, f) in enumerate([
    ("Streamlit UI", "app.py"),
    ("Conversational chat", "core/agents/chat.py"),
    ("MCP server (stdio)", "mcp_server.py"),
]):
    x = M + i * (sw + 0.44)
    shape(s, MSO_SHAPE.DOWN_ARROW, x + sw / 2 - 0.11, 2.74, 0.22, 0.26, STEEL)
    card(s, x, 3.08, sw, 0.92)
    text(s, n, x, 3.16, sw, 0.34, size=13.5, bold=True, align=PP_ALIGN.CENTER)
    text(s, f, x, 3.52, sw, 0.3, size=11, color=STEEL, align=PP_ALIGN.CENTER)
text(s, "A tool defined once cannot behave differently in chat than in MCP than in the graph.",
     M, 4.14, CW, 0.3, size=12.5, italic=True, color=MUTED, align=PP_ALIGN.CENTER)

rw = (CW - 0.66) / 4
for i, (k, v) in enumerate([
    ("Auth", "MCP runs as a local stdio process under the operator's own trust. The hosted "
             "REST profile in openapi.yaml specifies bearer tokens and per-tenant scoping."),
    ("Contracts", "Typed request/response per tool, mirrored in openapi.yaml and the MCP tool "
                  "schema, so a client sees one contract."),
    ("State", "Analysis is a job: submit, poll, fetch. Runs and evidence live in SQLite, so "
              "state outlives any one connection."),
    ("Multi-turn", "The client holds the conversation; the server holds run_id-addressable "
                   "facts. Follow-ups re-enter through the same tools."),
]):
    x = M + i * (rw + 0.22)
    card(s, x, 4.6, rw, 2.02, ICE_DEEP)
    text(s, k, x + 0.24, 4.78, rw - 0.48, 0.3, size=13.5, bold=True, color=AMBER)
    text(s, v, x + 0.24, 5.12, rw - 0.48, 1.4, size=11, color=NAVY)
notes(s, "Be explicit about the boundary: the MCP server is executable today; the REST "
         "connector is a specification, not a running service. Saying which is which is "
         "part of the deliverable.")

# --------------------------------------------------------- 10. KPI rationale
s = slide()
heading(s, "KPI rationale", "§3.4 — every metric has to be defensible")
text(s, "Category", M + 0.3, 1.62, 2.3, 0.28, size=11, bold=True, color=STEEL, spc=100)
text(s, "What it measures", M + 2.85, 1.62, 5.2, 0.28, size=11, bold=True, color=STEEL, spc=100)
text(s, "What it does NOT prove", M + 8.25, 1.62, 3.6, 0.28, size=11, bold=True, color=AMBER, spc=100)
y = 1.98
for i, (cat, what, isnot) in enumerate([
    ("Operational", "Latency p50/p95, cost per document, NOT_FOUND rate", "Throughput and spend — not quality"),
    ("Evidence", "Failed quote verifications / evidence-bearing attempts", "Source provenance — not entailment"),
    ("Follow-up", "Checks accepted after follow-up / checks sent to follow-up", "Agent-loop yield — not correctness"),
    ("Labelled evaluation", "State accuracy, confidence-calibration bins", "Only with a reviewed gold set"),
    ("Distribution", "Week-over-week state-share shift, by question", "Contract mix or regression — not accuracy"),
]):
    if i % 2 == 0:
        card(s, M, y, CW, 0.66)
    text(s, cat, M + 0.3, y, 2.45, 0.66, size=12.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, what, M + 2.85, y, 5.3, 0.66, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    text(s, isnot, M + 8.25, y, 3.75, 0.66, size=12, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.7

card(s, M, 6.24, CW, 0.8, NAVY)
text(s, "Provisional engineering thresholds, not measured SLOs:  quote-verification failures "
        "warn 5% / page 15%  ·  review rate > 20%  ·  budget exhaustion > 10%  ·  "
        "accuracy < 80%  ·  calibration gap > 20pp  ·  state drift > 25pp",
     M + 0.3, 6.24, CW - 0.6, 0.8, size=11.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Expect to defend each one. The right answer to 'what is your accuracy' is that no "
         "labelled ground truth ships with the assignment, so the dashboard renders accuracy "
         "as unavailable until a reviewed gold set exists — and an unavailable metric never "
         "fires an alert or borrows an older run's score.")

# ------------------------------------------------------- 11. observability
s = slide()
heading(s, "Monitoring and observability", "§3.5 / §3.6")
cw = (CW - 0.52) / 3
for i, (head, body) in enumerate([
    ("Correlated traces",
     "Every request carries a trace ID; every agent node opens a span. Router decisions, tool "
     "calls, verification results and critic reviews are all reconstructable from the span "
     "tree — live, during the interview."),
    ("Structured logs",
     "JSONL plus OpenTelemetry from one instrumentation path, exported to a local Langfuse on "
     "loopback. Confidential evidence stays out of any hosted tracer."),
    ("Operational window",
     "A 24-hour view against a seven-day daily-median baseline, with per-run anomalies — "
     "incomplete runs, backend failures, slowest runs — surfaced by name, not buried in an "
     "average."),
]):
    x = M + i * (cw + 0.26)
    card(s, x, 1.78, cw, 3.34)
    numdot(s, x + 0.34, 2.08, str(i + 1), 0.4)
    text(s, head, x + 0.34, 2.66, cw - 0.68, 0.34, size=15, bold=True)
    text(s, body, x + 0.34, 3.06, cw - 0.68, 1.95, size=12, color=MUTED)

card(s, M, 5.4, CW, 1.3, ICE_DEEP)
text(s, "Unavailable is a first-class value", M + 0.34, 5.58, 6.0, 0.3, size=14, bold=True)
text(s, "A missing metric renders as “unavailable”, never as 0.0, and never triggers an "
        "alert. Unknown model pricing and unaccounted failed calls stay explicit rather than "
        "being silently counted as zero — a dashboard that quietly reports zero is worse than "
        "one that admits it does not know.",
     M + 0.34, 5.92, CW - 0.68, 0.68, size=12.5, color=MUTED)
notes(s, "Offer to open a trace live. That is the fastest way to show routing and evaluation "
         "decisions rather than assert them.")

# ------------------------------------------------ 12. models, hosting, cost
s = slide()
heading(s, "Models, hosting and cost", "§4 / §10 — documented, not assumed")
y = 1.78
for i, (role, model, hosting) in enumerate([
    ("Inference", "qwen2.5:14b-instruct", "Local Ollama · localhost:11434/v1"),
    ("Embeddings", "nomic-embed-text, or FastEmbed BAAI/bge-small-en-v1.5", "Local, no API key"),
    ("Vector store", "Chroma", "Local, .index/chroma"),
    ("Orchestration", "LangGraph 1.2.11", "In-process"),
]):
    card(s, M, y, 7.9, 0.84, ICE if i % 2 == 0 else ICE_DEEP)
    text(s, role, M + 0.32, y, 1.9, 0.84, size=13, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, model, M + 2.24, y, 3.4, 0.84, size=12, anchor=MSO_ANCHOR.MIDDLE)
    text(s, hosting, M + 5.68, y, 2.4, 0.84, size=11.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.96

card(s, M + 8.22, 1.78, CW - 8.22, 3.54, NAVY)
text(s, "Cost basis", M + 8.54, 2.02, 3.4, 0.3, size=14, bold=True, color=AMBER)
text(s, [
    "Per-run cost comes from a per-million-token price table in core/pricing.py, recorded "
    "with cost_basis stating it is a table, not an invoice.",
    "Local inference is $0 in API terms. That means no API charge — not zero hardware or "
    "energy cost — and the dashboard says exactly that.",
], M + 8.54, 2.46, 3.5, 2.7, size=12, color=ICE, space_after=11)

card(s, M, 5.6, CW, 1.02, ICE_DEEP)
text(s, "Privacy gate", M + 0.32, 5.76, 3.0, 0.3, size=13.5, bold=True)
text(s, "The pipeline refuses a non-local inference or embedding config unless remote use is "
        "authorised explicitly, per upload, keyed to that document's own hash — so consent "
        "cannot leak from one document to the next.",
     M + 0.32, 6.08, CW - 0.64, 0.5, size=12, color=MUTED)
notes(s, "Live verification during development also used hosted providers, but only on "
         "non-confidential synthetic fixtures — never the sample contract. Say so plainly "
         "if asked.")

# ----------------------------------------------------------- 13. trade-offs
s = slide()
heading(s, "Key trade-offs", "What was chosen, and what it costs")
cw, ch = (CW - 0.36) / 2, 2.16
for i, (head, pro, con) in enumerate([
    ("Local-first over hosted", "Confidential contracts never leave the machine.",
     "A 14B local model is weaker than a frontier API model, and slower."),
    ("Decomposition over one prompt", "Per-facet attribution and a reproducible state rule.",
     "More model calls per question; a larger budget to manage."),
    ("Runtime discovery over annotation", "Generalises to a contract numbered differently.",
     "Weaker on the one sample than a hand-tuned map would be."),
    ("Explicit unavailable over zero", "No metric implies an accuracy nobody measured.",
     "The dashboard looks emptier than a less honest one."),
]):
    x = M + (i % 2) * (cw + 0.36)
    y = 1.76 + (i // 2) * (ch + 0.3)
    card(s, x, y, cw, ch)
    text(s, head, x + 0.32, y + 0.24, cw - 0.64, 0.36, size=15, bold=True)
    shape(s, MSO_SHAPE.OVAL, x + 0.32, y + 0.79, 0.15, 0.15, GREEN)
    text(s, pro, x + 0.6, y + 0.7, cw - 0.94, 0.58, size=12.5)
    shape(s, MSO_SHAPE.OVAL, x + 0.32, y + 1.41, 0.15, 0.15, RED)
    text(s, con, x + 0.6, y + 1.32, cw - 0.94, 0.62, size=12.5, color=MUTED)
notes(s, "Each of these I would defend but would also revisit under different constraints. "
         "The one I would not revisit is the last: a compliance tool that overstates what it "
         "measured is worse than no tool.")

# --------------------------------------------------- 14. gaps and next steps
s = slide(NAVY)
shape(s, MSO_SHAPE.OVAL, 11.2, -1.2, 3.8, 3.8, NAVY_SOFT)
text(s, "HONEST STATUS", M, 0.7, CW, 0.3, size=11.5, bold=True, color=AMBER, spc=200)
text(s, "Known gaps, and what comes next", M, 1.04, 11.0, 0.64, size=31, bold=True,
     color=WHITE, font=HEAD_FONT)
y = 2.06
for i, (head, body) in enumerate([
    ("Accuracy is unmeasured",
     "No labelled ground truth ships and there is one sample contract. The first real task is "
     "a reviewed gold set — until then the dashboard says unavailable, and so do I."),
    ("Acceptance run outstanding",
     "A reproducible real-model end-to-end run on the current configuration. Prior successful "
     "runs used independent non-confidential fixtures."),
    ("Scanned PDFs",
     "Text-extractable documents only. No OCR, and image content is not interpreted — image "
     "extraction degrades to a no-op rather than failing the pipeline."),
    ("REST connector",
     "openapi.yaml is a specification with an auth model, not a running service. MCP is the "
     "executable surface today."),
]):
    numdot(s, M, y + 0.02, str(i + 1), 0.38)
    text(s, head, M + 0.58, y, 11.4, 0.3, size=15, bold=True, color=WHITE)
    text(s, body, M + 0.58, y + 0.33, 11.4, 0.64, size=12.5, color=MUTED_LIGHT)
    y += 1.12

shape(s, MSO_SHAPE.RECTANGLE, M, 6.62, 1.1, 0.035, AMBER)
text(s, "Happy to open a live trace, the KPI tab, or any of this code.",
     M, 6.82, 11.0, 0.32, size=13.5, italic=True, color=WHITE)
notes(s, "Close on the gaps deliberately. Naming what is not done, and why, is more persuasive "
         "in a compliance context than claiming completeness.")

# ------------------------------------------------------------------ write
out = Path(__file__).resolve().parents[1] / "docs" / "walkthrough-deck.pptx"
prs.save(out)
print(f"wrote {out} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
