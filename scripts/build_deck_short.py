"""Build the 5-slide walkthrough deck (docs/walkthrough-deck-5.pptx).

A condensed cut of scripts/build_deck.py for a tight slot. It still covers every
topic the assignment's section 5.6 names -- architecture, agentic design,
connector/MCP design, KPI rationale, monitoring strategy, key trade-offs -- but
folds them into five slides, with the architecture drawn as a diagram rather
than described in prose.

Style helpers are duplicated from build_deck.py rather than imported, because
that module builds its 14 slides at import time.

    ./.venv/bin/python scripts/build_deck_short.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x16, 0x24, 0x3F)
NAVY_SOFT = RGBColor(0x24, 0x36, 0x5C)
STEEL = RGBColor(0x4A, 0x7F, 0xB5)
ICE = RGBColor(0xE8, 0xEE, 0xF7)
ICE_DEEP = RGBColor(0xD2, 0xDE, 0xEF)
AMBER = RGBColor(0xD9, 0x8A, 0x21)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5B, 0x6B, 0x85)
MUTED_LIGHT = RGBColor(0xA9, 0xB6, 0xCC)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

W, H = 13.333, 7.5
M = 0.62
CW = W - M * 2

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return s


def shape(s, kind, x, y, w, h, fill=None, adj=None):
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    if adj is not None:
        try:
            sp.adjustments[0] = adj
        except (IndexError, ValueError):
            pass
    sp.text_frame.text = ""
    return sp


def text(s, body, x, y, w, h, *, size=13, bold=False, italic=False, color=NAVY,
         font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None,
         line=None):
    box_ = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box_.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = body if isinstance(body, (list, tuple)) else [body]
    for i, entry in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line is not None:
            p.line_spacing = line
        run = p.add_run()
        run.text = entry
        f = run.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = color
        f.name = font
        if spacing is not None:
            run._r.get_or_add_rPr().set("spc", str(int(spacing)))
    return box_


def heading(s, title, kicker=None):
    if kicker:
        text(s, kicker.upper(), M, 0.46, CW, 0.26, size=11.5, bold=True,
             color=STEEL, spacing=90)
    text(s, title, M, 0.74, CW, 0.62, size=27, bold=True, color=NAVY, font=HEAD_FONT)
    shape(s, MSO_SHAPE.RECTANGLE, M, 1.44, 1.05, 0.035, AMBER)


def notes(s, body):
    s.notes_slide.notes_text_frame.text = body


def card(s, x, y, w, h, title, lines, *, fill=ICE, title_color=NAVY, body_color=MUTED,
         title_size=12.5, body_size=10.5):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, adj=0.07)
    text(s, title, x + 0.16, y + 0.11, w - 0.32, 0.26, size=title_size, bold=True,
         color=title_color)
    if lines:
        text(s, lines, x + 0.16, y + 0.42, w - 0.32, h - 0.54, size=body_size,
             color=body_color, line=1.18)


def down(s, x, y):
    shape(s, MSO_SHAPE.DOWN_ARROW, x, y, 0.24, 0.26, STEEL)


# ---------------------------------------------------------------- 1 · title
s = slide(NAVY)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, 0.10, AMBER)
text(s, "CONTRACT ANALYZER", M, 2.10, CW, 0.92, size=44, bold=True, color=WHITE,
     font=HEAD_FONT, spacing=40)
text(s, "Agentic RAG compliance analysis over third-party contracts",
     M, 3.10, CW, 0.44, size=19, color=MUTED_LIGHT)
shape(s, MSO_SHAPE.RECTANGLE, M, 3.80, 1.05, 0.035, AMBER)

x = M
for label, detail in [
    ("Upload", "any text-extractable contract PDF"),
    ("Analyse", "5 questions → 25 evidence checks"),
    ("Return", "validated JSON, verbatim cited quotes"),
]:
    card(s, x, 4.30, 3.82, 1.02, label, [detail], fill=NAVY_SOFT,
         title_color=AMBER, body_color=MUTED_LIGHT, title_size=13, body_size=11.5)
    x += 4.02

text(s, "Architecture · agentic design · connector and MCP · KPIs and monitoring · trade-offs",
     M, 6.40, CW, 0.32, size=13, italic=True, color=MUTED_LIGHT)
notes(s, "Five slides. The constraint behind every decision: this has to work on a contract "
         "nobody has seen, not just the sample that shipped with the assignment.")

# --------------------------------------------------------- 2 · architecture
s = slide()
heading(s, "One service layer, one tool registry", "Architecture")

card(s, M, 1.64, 5.90, 0.78, "Streamlit UI  ·  app.py, ui/",
     ["upload · results · chat · KPI · monitoring"], fill=ICE_DEEP)
card(s, M + 6.20, 1.64, 5.89, 0.78, "MCP client  ·  mcp_server.py",
     ["stdio JSON-RPC, no tool of its own"], fill=ICE_DEEP)
down(s, 3.45, 2.48)
down(s, 9.65, 2.48)

card(s, M, 2.84, CW, 0.90, "Shared service  ·  core/",
     ["pipeline.analyze() — the single entry point        "
      "ToolRegistry — one definition, three consumers        agents/chat.py"],
     fill=STEEL, title_color=WHITE, body_color=WHITE, body_size=11)
down(s, 6.55, 3.80)

x, cw = M, (CW - 0.30) / 4
for title, detail in [
    ("ingest/", "layout PDF → blocks,\ntables, structural chunks"),
    ("retrieval/", "lexical + vector RRF,\nparent and reference reads"),
    ("rubric/", "25 checks with facets,\nno contract IDs inside"),
    ("agents/graph.py", "LangGraph bounded\nevidence loop"),
]:
    card(s, x, 4.18, cw, 1.10, title, detail.split("\n"), fill=ICE)
    x += cw + 0.10

card(s, M, 5.48, 6.00, 0.82, "SQLite  ·  core/store/",
     ["runs · findings · evidence snapshots · ordered audit"], fill=ICE)
card(s, M + 6.30, 5.48, 5.79, 0.82, "Observability  ·  core/obs/",
     ["JSON logs + OTel · trace_id / span_id / parent_span_id"], fill=ICE)

text(s, "The core/ boundary is load-bearing: UI, chat and MCP import the same functions, so the "
        "MCP deliverable is the real system rather than a parallel demo of it.",
     M, 6.54, CW, 0.40, size=12.5, italic=True, color=MUTED, line=1.20)
notes(s, "Read it top down. Two clients, one service layer, one registry — that is the point. "
         "A tool defined once cannot behave differently in chat than in MCP than in the graph. "
         "The rubric is data, not code, so the checks change without touching orchestration. "
         "REST is a specification only, which is why it is deliberately not on this diagram.")

# ------------------------------------------------------- 3 · agentic design
s = slide()
heading(s, "Decompose, retrieve, judge, verify, critique", "Agentic design  ·  §3.2")

x, cw = M, (CW - 4 * 0.40) / 5
nodes = [("Router", "plans tools per\npending check"),
         ("Retrieve", "search · context ·\nreferences · modifiers"),
         ("Judge", "every facet, with\nevidence indices"),
         ("Verify", "mechanical —\ncannot be overruled"),
         ("Critic", "entailment, scope,\ndates, exceptions")]
for i, (title, detail) in enumerate(nodes):
    card(s, x, 1.68, cw, 1.16, title, detail.split("\n"),
         fill=ICE_DEEP if title in {"Verify", "Critic"} else ICE)
    if i < 4:
        shape(s, MSO_SHAPE.RIGHT_ARROW, x + cw + 0.05, 2.14, 0.30, 0.24, STEEL)
    x += cw + 0.40

shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, M, 3.02, CW, 0.42, ICE, adj=0.30)
text(s, "targeted evidence gap → back to Router     ·     bounded: 2 follow-ups, 32 LLM calls, "
        "120k tokens, 48k evidence characters per check",
     M + 0.24, 3.11, CW - 0.5, 0.28, size=11.5, color=NAVY)

for x0, title, body in [
    (M, "Decomposition is structural, not prompted",
     "25 checks with explicit required facets, grouped 7/4/4/5/5. No section numbers, exhibit "
     "letters or control-ID prefixes anywhere in the rubric — a table transcribed from one PDF "
     "annotates that PDF, it does not analyse contracts."),
    (M + 6.30, "Verification is mechanical, then semantic",
     "The quote must be a substring of the cited chunk and of the document; NOT_FOUND may carry "
     "no quote; an applied modifier needs source evidence. The critic may ask for more, but "
     "cannot overrule the machine. Aggregation to the three states is a rule, not a model."),
]:
    text(s, title, x0, 3.76, 5.80, 0.30, size=14, bold=True, color=NAVY, font=HEAD_FONT)
    text(s, body, x0, 4.14, 5.80, 1.50, size=12.5, color=MUTED, line=1.30)

shape(s, MSO_SHAPE.RECTANGLE, M, 5.86, 1.05, 0.035, AMBER)
text(s, "Failure never becomes a verdict — a backend error is ERROR, malformed output is "
        "UNRESOLVED, and both surface as review items with the run marked incomplete.",
     M, 6.06, CW, 0.60, size=13, color=NAVY, line=1.25)
notes(s, "The two shaded nodes are the evaluator the assignment asks for, deliberately split. "
         "Mechanical verification is the hallucination detector and it is not negotiable. The "
         "critic is judgement and it is advisory. Confidence is the minimum across both and is "
         "documented as a heuristic — I am not claiming calibration I have not measured.")

# ------------------------------------------- 4 · connector, KPIs, monitoring
s = slide()
heading(s, "Exposed, measured, observable", "Connector · KPIs · monitoring  ·  §3.3–§3.6")

card(s, M, 1.66, 3.87, 2.24, "Connector / MCP",
     ["MCP is executable: the stdio server defines no tool of its own, it dispatches "
      "onto the shared registry.",
      "REST is an OpenAPI spec, not a running server — per-tenant API key, run_id and "
      "session_id state, 202-then-poll because a tool call must return promptly."],
     body_size=11)
card(s, M + 4.07, 1.66, 3.87, 2.24, "KPI rationale",
     ["Six categories kept apart, because collapsing them is how a dashboard starts "
      "implying accuracy it never measured:",
      "operations · reliability · evidence provenance · follow-up yield · gold-set "
      "quality · LLM-judged.",
      "Thresholds are engineering guardrails, not SLOs."], body_size=11)
card(s, M + 8.14, 1.66, 3.87, 2.24, "Monitoring",
     ["Structured JSON to stderr and file — never stdout, because MCP speaks JSON-RPC there.",
      "trace_id / span_id / parent_span_id on every record, so one request reconstructs "
      "parse → route → judge → verify → critic.",
      "Local Langfuse over the same OTel tree."], body_size=11)

text(s, ["Metric", "Threshold", "What it does and does not prove"], M + 0.14, 4.16, CW, 0.26,
     size=11, bold=True, color=STEEL, spacing=60)
y = 4.52
for i, (name, thr, meaning) in enumerate([
    ("Quote verification failure rate", "warn 5%  ·  page 15%", "fabricated citations, not entailment"),
    ("Checks needing review", "warn 20%", "analysis availability, not contract compliance"),
    ("Budget exhaustion", "warn 10%", "loop cost, before it shows up as latency"),
    ("Cost per document", "2× the 7-day median", "the number finance asks about"),
]):
    if i % 2 == 0:
        shape(s, MSO_SHAPE.RECTANGLE, M, y - 0.07, CW, 0.46, ICE)
    text(s, name, M + 0.14, y, 4.40, 0.28, size=12, bold=True, color=NAVY)
    text(s, thr, M + 4.80, y, 2.70, 0.28, size=12, bold=True, color=AMBER)
    text(s, meaning, M + 7.60, y, 4.30, 0.28, size=12, color=MUTED)
    y += 0.52

notes(s, "Three requirements on one slide because they are one story: exposed through a shared "
         "registry, measured by metrics kept honest about what they do not prove, traceable end "
         "to end. Every KPI's definition, formula, threshold and action live in one structure, "
         "so the dashboard and the MCP explain_kpi tool cannot drift apart.")

# --------------------------------------------------------- 5 · trade-offs
s = slide(NAVY)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, 0.10, AMBER)
text(s, "TRADE-OFFS", M, 0.50, CW, 0.28, size=11.5, bold=True, color=STEEL, spacing=90)
text(s, "What was chosen, and what it costs", M, 0.84, CW, 0.56, size=27, bold=True,
     color=WHITE, font=HEAD_FONT)

y = 1.68
for title, detail in [
    ("Layout-first parsing",
     "Tables carry answers prose does not — one of the five questions is answerable only from "
     "an exhibit row. Cost: complex multi-column layouts stay heuristic, and are flagged rather "
     "than silently accepted."),
    ("Decomposition plus deterministic aggregation",
     "25 checks instead of one prompt per question. Cost: more model calls. Gain: the verdict is "
     "auditable facet by facet, and the compliance state is a rule rather than an opinion."),
    ("Local-first inference",
     "Confidential material never leaves the machine; remote inference needs explicit "
     "per-document consent. Cost: throughput, and a free tier that rate-limits."),
]:
    shape(s, MSO_SHAPE.RECTANGLE, M, y + 0.04, 0.035, 0.62, AMBER)
    text(s, title, M + 0.24, y, 11.60, 0.30, size=15, bold=True, color=WHITE, font=HEAD_FONT)
    text(s, detail, M + 0.24, y + 0.34, 11.60, 0.52, size=12.5, color=MUTED_LIGHT, line=1.22)
    y += 1.04

shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, M, 4.96, CW, 1.44, NAVY_SOFT, adj=0.08)
text(s, "Honest status", M + 0.26, 5.10, 11.60, 0.28, size=14, bold=True, color=AMBER,
     font=HEAD_FONT)
text(s, ["The five questions return a compliance state on any text-extractable contract, with "
         "verbatim quotes and document-derived references — verified on a PDF that is not the "
         "sample.",
         "No labelled ground truth ships with this assignment and there is one sample contract, "
         "so no accuracy figure is claimed. Open gaps are listed in the README, not hidden."],
     M + 0.26, 5.44, 11.60, 0.84, size=12.5, color=WHITE, line=1.24)

text(s, "Happy to open a live trace, the KPI tab, or any of this code.",
     M, 6.70, CW, 0.32, size=13.5, italic=True, color=MUTED_LIGHT)
notes(s, "Close on the trade-offs, then on what is not done. In a compliance context, naming the "
         "gaps is more persuasive than claiming completeness — the panel will trust the numbers "
         "I do give them more if I am precise about the ones I cannot.")

out = Path(__file__).resolve().parents[1] / "docs" / "walkthrough-deck-5.pptx"
prs.save(out)
print(f"wrote {out} ({len(prs.slides._sldIdLst)} slides)")
